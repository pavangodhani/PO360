"""Safe Excel report generation.

The database uses sqlite3.Row objects. A sqlite Row raises IndexError when a
column requested by name does not exist. The previous exporter assumed every
history column existed exactly as named, so an otherwise successful mailbox
run could fail while generating the workbook.

This replacement uses safe row access and column aliases. Missing optional
columns become blank instead of crashing the entire report.
"""

from __future__ import annotations

import json
import logging
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

LOGGER = logging.getLogger(__name__)

MASTER_COLUMNS = [
    "S No", "Company", "Order By", "Order Date", "Address Date",
    "State / Region", "Branch Name", "Cluster / MC Name", "MC Code",
    "Address", "BH / CBH Name", "BH / CBH Contact No", "Goods Required",
    "Quantity", "PO Number", "PO Status", "Current Status",
    "Goods Dispatch Status", "Delivery Status", "Mail Received Status",
    "Invoice Send Status", "Done By", "Source Mailbox",
    "Last Updated Date", "Remarks",
]


def _safe(row: Any, *keys: str, default: str = "") -> Any:
    """Read a sqlite Row/dict/object without raising for missing columns."""
    if row is None:
        return default

    for key in keys:
        try:
            if isinstance(row, dict):
                value = row.get(key)
            else:
                # sqlite3.Row supports key lookup but raises IndexError when
                # the key isn't present.
                try:
                    value = row[key]
                except (KeyError, IndexError):
                    value = None

            if value is not None:
                return value
        except (KeyError, IndexError, TypeError):
            continue

    return default


def _event_details(raw: Any) -> str:
    if raw is None:
        return ""

    if isinstance(raw, dict):
        return str(raw.get("value") or raw.get("evidence") or "")

    text = str(raw)
    if not text:
        return ""

    try:
        obj = json.loads(text)
        if isinstance(obj, dict):
            return str(obj.get("value") or obj.get("evidence") or obj.get("details") or "")
    except (TypeError, ValueError, json.JSONDecodeError):
        pass

    return text


def generate_report(db, output_path: str | Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "PO Report"
    ws.append(MASTER_COLUMNS)

    rows = db.get_report_rows()

    for index, row in enumerate(rows, start=1):
        values = [
            index,
            _safe(row, "company"),
            _safe(row, "order_by"),
            _safe(row, "order_date"),
            _safe(row, "address_date"),
            _safe(row, "state", "state_region"),
            _safe(row, "branch_name"),
            _safe(row, "cluster_name", "mc_name"),
            _safe(row, "mc_code"),
            _safe(row, "address"),
            _safe(row, "bh_name", "cbh_name"),
            _safe(row, "bh_contact", "cbh_contact"),
            _safe(row, "goods_required", "goods"),
            _safe(row, "quantities", "quantity"),
            _safe(row, "po_number"),
            _safe(row, "po_status"),
            _safe(row, "current_status", "status"),
            _event_details(_safe(row, "dispatch_status_event", "goods_dispatch_status")),
            _event_details(_safe(row, "delivery_status_event", "delivery_status")),
            _event_details(_safe(row, "mail_status_event", "mail_received_status")),
            _event_details(_safe(row, "invoice_status_event", "invoice_status")),
            _safe(row, "done_by"),
            _safe(row, "source_mailbox", "mailbox"),
            _safe(row, "last_event_date", "updated_at", "last_updated_date"),
            _safe(row, "remarks"),
        ]
        ws.append(values)

    _style_sheet(ws, freeze="A2", autofilter=True)

    history = wb.create_sheet("PO History")
    history_columns = [
        "PO Number", "Event Type", "Event Date", "Details",
        "Email Subject", "Sender", "Received At",
    ]
    history.append(history_columns)

    history_rows = db.get_history_rows()
    for row in history_rows:
        # Do NOT use row[c] directly: some database versions may call these
        # columns event_date/created_at or subject/from_address.
        history.append([
            _safe(row, "po_number"),
            _safe(row, "event_type"),
            _safe(row, "event_date", "created_at", "received_at"),
            _event_details(_safe(row, "details", "event_details", "value")),
            _safe(row, "email_subject", "subject"),
            _safe(row, "sender", "from_address"),
            _safe(row, "received_at", "email_received_at", "event_date"),
        ])

    _style_sheet(history, freeze="A2", autofilter=True)

    errors = wb.create_sheet("Processing Errors")
    error_columns = [
        "Mailbox", "Message ID", "Received At", "Subject", "Status", "Error",
    ]
    errors.append(error_columns)

    for row in db.get_error_rows():
        errors.append([
            _safe(row, "mailbox", "source_mailbox"),
            _safe(row, "message_id"),
            _safe(row, "received_at", "created_at"),
            _safe(row, "subject", "email_subject"),
            _safe(row, "status"),
            _safe(row, "error", "error_message", "details"),
        ])

    _style_sheet(errors, freeze="A2", autofilter=True)

    try:
        wb.save(output_path)
    except PermissionError as exc:
        LOGGER.error(
            "Cannot write Excel report '%s'. It may be open in Excel.",
            output_path,
            exc_info=True,
        )
        raise RuntimeError(
            f"Cannot write {output_path}. Close the workbook in Excel and run again."
        ) from exc

    LOGGER.info(
        "Excel report generated: %s (%s PO rows)",
        output_path,
        len(rows),
    )
    return output_path


def _style_sheet(ws, freeze: str, autofilter: bool) -> None:
    ws.freeze_panes = freeze

    if autofilter and ws.max_row >= 1:
        ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for col_cells in ws.columns:
        max_len = 0
        for cell in col_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(value), 60))

        width = max(12, min(max_len + 2, 60))
        ws.column_dimensions[
            get_column_letter(col_cells[0].column)
        ].width = width

    ws.row_dimensions[1].height = 32